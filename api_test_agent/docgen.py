"""Test-case document generation (xlsx + markdown) and user-review readback."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_COLUMNS = ["用例编号", "所属模块", "用例标题", "优先级", "设计方法", "前置条件",
                   "测试步骤", "测试数据", "预期结果", "实际结果", "执行状态", "关联自动化用例"]
# 用户评审时允许编辑并回读的列
EDITABLE = {"用例编号": "case_id", "所属模块": "module", "用例标题": "title",
            "优先级": "priority", "设计方法": "design_method", "前置条件": "precondition",
            "测试步骤": "steps", "测试数据": "test_data", "预期结果": "expected"}
COLUMN_WIDTHS = {"用例编号": 16, "所属模块": 10, "用例标题": 34, "优先级": 8, "设计方法": 10,
                 "前置条件": 22, "测试步骤": 36, "测试数据": 34, "预期结果": 34,
                 "实际结果": 30, "执行状态": 10, "关联自动化用例": 34}


def _col_width(columns: list[str]) -> dict:
    return {c: COLUMN_WIDTHS.get(c, 16) for c in columns}


def _cell_value(value):
    """openpyxl 单元格只收标量;LLM 可能把步骤等字段写成数组/对象,统一转文本。"""
    if isinstance(value, (list, dict)):
        import json
        return json.dumps(value, ensure_ascii=False)
    return value


# ---------------------------------------------------------------- xlsx
def write_excel(cases: list[dict], path: str | Path, columns: list[str] | None = None,
                overview: dict | None = None, endpoints: list[dict] | None = None) -> Path:
    columns = columns or DEFAULT_COLUMNS
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "测试用例"
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    ws.append(columns)
    for col in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = _col_width(columns)[columns[col - 1]]
    ws.freeze_panes = "A2"

    for case in cases:
        row = []
        for col_name in columns:
            field = EDITABLE.get(col_name)
            if field:
                row.append(_cell_value(case.get(field, "")))
            elif col_name == "实际结果":
                row.append(_cell_value(case.get("actual_result", "")))
            elif col_name == "执行状态":
                row.append(_cell_value(case.get("exec_status", "未执行")))
            elif col_name == "关联自动化用例":
                row.append(_cell_value(case.get("pytest_node", "")))
            else:
                row.append("")
        ws.append(row)
        for col in range(1, len(columns) + 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.border = border
            cell.alignment = wrap

    _add_overview_sheet(wb, cases, overview or {}, endpoints or [])
    wb.save(path)
    return path


def _add_overview_sheet(wb: Workbook, cases: list[dict], overview: dict, endpoints: list[dict]) -> None:
    ws = wb.create_sheet("测试概览")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    rows = [
        ("项目", overview.get("project", "")),
        ("被测目录", overview.get("project_dir", "")),
        ("需求来源", overview.get("requirement_source", "全量模式")),
        ("接口数量", len(endpoints)),
        ("用例总数", len(cases)),
        ("生成时间", overview.get("generated_at", "")),
        ("设计 Skill", overview.get("design_skill", "")),
        ("文档 Skill", overview.get("doc_skill", "")),
    ]
    for key, value in rows:
        ws.append([key, str(value)])
    ws.append([])

    ws.append(["按优先级", "数量"])
    for p in ("P0", "P1", "P2", "P3"):
        ws.append([p, sum(1 for c in cases if c.get("priority") == p)])
    ws.append([])
    ws.append(["按设计方法", "数量"])
    for m in ("等价类", "边界值", "场景法", "因果图"):
        ws.append([m, sum(1 for c in cases if c.get("design_method") == m)])

    if endpoints:
        ws2 = wb.create_sheet("接口清单")
        ws2.append(["方法", "路径", "模块", "摘要", "需鉴权"])
        for col, width in zip("ABCDE", (8, 34, 12, 44, 8)):
            ws2.column_dimensions[col].width = width
        for ep in endpoints:
            ws2.append([ep.get("method", ""), ep.get("path", ""), ep.get("module", ""),
                        ep.get("summary", ""), "是" if ep.get("auth_required") else "否"])


# ---------------------------------------------------------------- markdown
def _md_escape(value) -> str:
    text = str(value if value is not None else "")
    return text.replace("|", "\\|").replace("\n", "<br>")


def write_markdown(cases: list[dict], path: str | Path, columns: list[str] | None = None,
                   overview: dict | None = None, endpoints: list[dict] | None = None) -> Path:
    columns = columns or DEFAULT_COLUMNS
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    ov = overview or {}
    lines = [
        f"# 测试用例 — {ov.get('project', '')}",
        "",
        f"- 被测目录:`{ov.get('project_dir', '')}`",
        f"- 需求来源:{ov.get('requirement_source', '全量模式')}",
        f"- 接口数量:{len(endpoints or [])};用例总数:**{len(cases)}**"
        f"(P0:{sum(1 for c in cases if c.get('priority') == 'P0')},"
        f"P1:{sum(1 for c in cases if c.get('priority') == 'P1')},"
        f"P2:{sum(1 for c in cases if c.get('priority') == 'P2')},"
        f"P3:{sum(1 for c in cases if c.get('priority') == 'P3')})",
        f"- 设计 Skill:{ov.get('design_skill', '')};文档 Skill:{ov.get('doc_skill', '')}",
        "",
        "> 评审说明:可直接修改本表或同目录 Excel(两文件同源);"
        "修改后执行确认,Agent 将以修改后的内容为基线生成自动化脚本。",
        "",
    ]
    modules: list[str] = []
    for case in cases:
        module = case.get("module", "COMMON")
        if module not in modules:
            modules.append(module)
    for module in modules:
        lines.append(f"## 模块:{module}")
        lines.append("")
        lines.append("| " + " | ".join(columns) + " |")
        lines.append("|" + "|".join(["---"] * len(columns)) + "|")
        for case in [c for c in cases if c.get("module", "COMMON") == module]:
            row = []
            for col_name in columns:
                field = EDITABLE.get(col_name)
                if field:
                    row.append(case.get(field, ""))
                elif col_name == "实际结果":
                    row.append(case.get("actual_result", ""))
                elif col_name == "执行状态":
                    row.append(case.get("exec_status", "未执行"))
                elif col_name == "关联自动化用例":
                    row.append(case.get("pytest_node", ""))
                else:
                    row.append("")
            lines.append("| " + " | ".join(_md_escape(v) for v in row) + " |")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ---------------------------------------------------------------- readback(评审回读)
def readback_cases(path: str | Path, existing: list[dict],
                   columns: list[str] | None = None) -> tuple[list[dict], list[str]]:
    """把用户评审后的文档读回为用例基线。

    以 case_id 为主键合并:用户改的字段覆盖原值;被用户删除的用例移除;
    用户新增的行作为手工用例保留(无 request,不参与脚本生成)。
    返回 (新用例列表, 警告列表)。
    """
    path = Path(path)
    columns = columns or DEFAULT_COLUMNS
    if path.suffix == ".xlsx":
        rows, warnings = _readback_excel(path, columns)
    else:
        rows, warnings = _readback_markdown(path, columns)

    by_id = {c["case_id"]: c for c in existing}
    result: list[dict] = []
    seen: set[str] = set()
    for row in rows:
        cid = str(row.get("case_id", "")).strip()
        if not cid:
            warnings.append(f"存在无编号行,已保留为手工用例:{row.get('title', '')!r}")
            row["manual"] = True
            row["exec_status"] = row.get("exec_status", "未执行")
            result.append(row)
            continue
        seen.add(cid)
        base = by_id.get(cid)
        if base is None:
            row["manual"] = True
            row["exec_status"] = row.get("exec_status", "未执行")
            result.append(row)
            continue
        merged = dict(base)
        for field in EDITABLE.values():
            new_value = row.get(field, "")
            if str(new_value).strip():
                merged[field] = new_value
        result.append(merged)
    for cid, case in by_id.items():
        if cid not in seen:
            warnings.append(f"用例 {cid} 在评审后文档中已被删除,不参与后续生成")
    return result, warnings


def _readback_excel(path: Path, columns: list[str]) -> tuple[list[dict], list[str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb["测试用例"] if "测试用例" in wb.sheetnames else wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col_index = {name: i for i, name in enumerate(header)}
    missing = [c for c in columns[:5] if c not in col_index]
    if missing:
        return [], [f"Excel 表头缺少关键列: {missing};请勿删除表头"]
    rows: list[dict] = []
    warnings: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        case = {field: _clean(row[col_index[col]]) for col, field in EDITABLE.items()
                if col in col_index}
        if "执行状态" in col_index:
            case["exec_status"] = _clean(row[col_index["执行状态"]]) or "未执行"
        rows.append(case)
    return rows, warnings


def _readback_markdown(path: Path, columns: list[str]) -> tuple[list[dict], list[str]]:
    text = path.read_text(encoding="utf-8")
    rows: list[dict] = []
    warnings: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|") or set(line) <= {"|", "-", " "}:
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells and cells[0] in ("用例编号", "---", ""):
            continue
        if len(cells) != len(columns):
            warnings.append(f"Markdown 行列数不符,已跳过:{line[:50]}...")
            continue
        case = {field: _clean_cells(cells[i]) for i, (col, field) in enumerate(EDITABLE.items())
                if col in columns}
        idx = columns.index("执行状态") if "执行状态" in columns else -1
        if idx >= 0:
            case["exec_status"] = cells[idx] or "未执行"
        rows.append(case)
    return rows, warnings


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clean_cells(value: str) -> str:
    return value.replace("<br>", "\n").replace("\\|", "|").strip()
