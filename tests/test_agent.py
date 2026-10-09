"""Unit tests for the api-test-agent core modules."""
import textwrap
from pathlib import Path

import pytest

from api_test_agent import docgen
from api_test_agent.pi_agent import text_tool_result
from api_test_agent.requirements_input import load_requirement
from api_test_agent.skills import (SKILL_TYPE_DESIGN, SKILL_TYPE_DOC,
                                   default_case_id_pattern, discover_skills,
                                   parse_skill_md)
from api_test_agent.state import SessionState
from api_test_agent.testframe.generator import generate_framework, write_data_files
from api_test_agent.validators import path_matches, validate_case, validate_module_cases


# ---------------------------------------------------------------- tool result shape
def test_tool_result_uses_content_block():
    """SDK 要求 content 为 TextContentBlock;dict 会导致结果回传失败。"""
    result = text_tool_result("ok")
    block = result.content[0]
    assert block.type == "text" and block.text == "ok"


# ---------------------------------------------------------------- skills
def test_builtin_skills_discovered():
    found, warnings = discover_skills("does-not-matter")
    assert not warnings
    assert SKILL_TYPE_DESIGN in found and SKILL_TYPE_DOC in found
    assert found[SKILL_TYPE_DESIGN].level == "builtin"
    assert "等价类" in found[SKILL_TYPE_DESIGN].body
    assert default_case_id_pattern(found[SKILL_TYPE_DOC]).startswith("^")


def test_project_skill_overrides_builtin(tmp_path: Path):
    custom = tmp_path / ".api-test-agent" / "skills" / "acme-doc-format"
    custom.mkdir(parents=True)
    (custom / "SKILL.md").write_text(textwrap.dedent("""\
        ---
        name: acme-doc-format
        type: doc-format
        description: ACME 模板
        case_id_pattern: "^ACME-[A-Z]+-\\\\d{4}$"
        ---
        # ACME 编号规则
        项目码-模块码-四位序号。
        """), encoding="utf-8")
    found, warnings = discover_skills(tmp_path)
    assert not warnings
    doc = found[SKILL_TYPE_DOC]
    assert doc.name == "acme-doc-format" and doc.level == "project"
    assert default_case_id_pattern(doc) == r"^ACME-[A-Z]+-\d{4}$"


def test_bad_skill_type_reported_not_fatal(tmp_path: Path):
    bad = tmp_path / ".api-test-agent" / "skills" / "broken"
    bad.mkdir(parents=True)
    (bad / "SKILL.md").write_text("---\nname: x\ntype: wrong\n---\nbody", encoding="utf-8")
    found, warnings = discover_skills(tmp_path)
    assert len(warnings) == 1
    assert found[SKILL_TYPE_DOC].level == "builtin"


def test_parse_skill_md_requires_front_matter(tmp_path: Path):
    skill_md = tmp_path / "SKILL.md"
    skill_md.write_text("no front matter", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_skill_md(skill_md, "builtin")


# ---------------------------------------------------------------- validators
def _case(**overrides):
    base = {
        "case_id": "DEMO-USER-001", "module": "USER",
        "title": "注册成功", "priority": "P0", "design_method": "等价类",
        "expected": "HTTP 200",
        "request": {"method": "POST", "path": "/api/register", "json": {"username": "a"}},
        "expected_spec": {"status_code": 200, "code": 0},
    }
    base.update(overrides)
    return base


def test_validate_case_ok():
    assert validate_case(_case(), 1, "USER", r"^DEMO-[A-Z]+-\d{3}$") == []


def test_validate_case_bad_priority_and_id():
    errors = validate_case(_case(priority="PX", case_id="wrong-1"), 1, "USER",
                           r"^DEMO-[A-Z]+-\d{3}$")
    assert any("priority" in e for e in errors)
    assert any("case_id" in e for e in errors)


def test_validate_case_scenario_skips_request():
    scenario_case = _case(request=None, expected_spec=None, scenario_steps=[
        {"method": "POST", "path": "/api/register", "expect_status": 200},
        {"method": "GET", "path": "/api/users/$step1.data.id", "expect_status": 200},
    ])
    assert validate_case(scenario_case, 1, "USER", r".+") == []
    bad_step = _case(request=None, expected_spec=None,
                     scenario_steps=[{"path": "/api/users"}])
    assert validate_case(bad_step, 1, "USER", r".+")


def test_path_matches_template():
    assert path_matches("/api/users/{user_id}", "/api/users/1")
    assert path_matches("/api/users", "/api/users")
    assert not path_matches("/api/users/{user_id}", "/api/users/1/orders")
    assert not path_matches("/api/users/{user_id}", "/api/orders/1")


def test_module_coverage_errors():
    endpoints = [{"method": "GET", "path": "/api/users", "module": "USER", "params": []}]
    cases = [_case(request={"method": "GET", "path": "/api/users", "query": {}})]
    errors = validate_module_cases(cases, "USER", r".+", endpoints)
    assert not any("没有任何用例" in e for e in errors)
    # 缺 P0
    cases[0]["priority"] = "P1"
    errors = validate_module_cases(cases, "USER", r".+", endpoints)
    assert any("P0" in e for e in errors)


# ---------------------------------------------------------------- docgen
def _sample_cases():
    return [
        {"case_id": "DEMO-AUTH-001", "module": "AUTH", "title": "登录成功",
         "priority": "P0", "design_method": "等价类", "precondition": "有账号",
         "steps": "1. POST /api/login", "test_data": '{"username":"admin"}',
         "expected": "code=0", "request": {"method": "POST", "path": "/api/login",
                                           "json": {"username": "admin"}, "auth": False},
         "expected_spec": {"status_code": 200, "code": 0}},
        {"case_id": "DEMO-AUTH-002", "module": "AUTH", "title": "登录失败",
         "priority": "P1", "design_method": "等价类", "precondition": "",
         "steps": "1. POST /api/login", "test_data": "{}", "expected": "401",
         "request": {"method": "POST", "path": "/api/login", "json": {}, "auth": False},
         "expected_spec": {"status_code": 401, "code": 401}},
    ]


def test_excel_roundtrip_with_user_edit(tmp_path: Path):
    cases = _sample_cases()
    xlsx = docgen.write_excel(cases, tmp_path / "cases.xlsx")
    md = docgen.write_markdown(cases, tmp_path / "cases.md")
    assert xlsx.exists() and md.exists()

    # 模拟用户:改标题、删第二条、加一条新用例
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["测试用例"]
    ws.cell(row=2, column=3, value="登录成功(改)")
    ws.delete_rows(3)
    ws.append(["DEMO-AUTH-009", "AUTH", "手工补充", "P2", "等价类", "", "1. x", "", "code=0"])
    wb.save(xlsx)

    merged, warnings = docgen.readback_cases(xlsx, cases)
    by_id = {c["case_id"]: c for c in merged}
    assert by_id["DEMO-AUTH-001"]["title"] == "登录成功(改)"
    assert "DEMO-AUTH-002" not in by_id
    assert by_id["DEMO-AUTH-009"]["manual"] is True
    assert any("已被删除" in w for w in warnings)


def test_markdown_roundtrip(tmp_path: Path):
    cases = _sample_cases()
    md = docgen.write_markdown(cases, tmp_path / "cases.md")
    merged, warnings = docgen.readback_cases(md, cases)
    assert not warnings
    assert [c["case_id"] for c in merged] == ["DEMO-AUTH-001", "DEMO-AUTH-002"]


# ---------------------------------------------------------------- requirements input
def test_requirement_inline_text():
    doc = load_requirement(text="支持商品条件筛选")
    assert doc.source == "inline-text" and "商品" in doc.text


def test_requirement_txt_file(tmp_path: Path):
    req = tmp_path / "req.txt"
    req.write_text("本版本新增订单导出", encoding="utf-8")
    assert "订单导出" in load_requirement(file=str(req)).text


def test_requirement_docx_unsupported(tmp_path: Path):
    req = tmp_path / "req.docx"
    req.write_text("x", encoding="utf-8")
    with pytest.raises(NotImplementedError):
        load_requirement(file=str(req))


def test_requirement_empty_raises():
    with pytest.raises(ValueError):
        load_requirement(text="  ")


# ---------------------------------------------------------------- state
def test_state_roundtrip(tmp_path: Path):
    state = SessionState()
    state.endpoints = [{"method": "GET", "path": "/x", "module": "X"}]
    state.cases = [{"case_id": "X-1", "module": "X"}]
    file = tmp_path / "state.json"
    state.save(file)
    loaded = SessionState.load(file)
    assert loaded.session_id == state.session_id
    assert loaded.endpoints == state.endpoints
    assert loaded.next_case_id("X", "DEMO") == "DEMO-X-002"


# ---------------------------------------------------------------- codegen
def test_generate_framework(tmp_path: Path):
    class FakeService:
        base_url = "http://127.0.0.1:8100"
        timeout = 15
        app_cmd = ""
        test_account = {"username": "admin", "password": "x"}
        login = {"path": "/api/login"}

    class FakeCfg:
        project_code = "DEMO"
        service = FakeService()

    cases = _sample_cases() + [{
        "case_id": "DEMO-AUTH-003", "module": "AUTH", "title": "场景链", "priority": "P1",
        "design_method": "场景法", "expected": "ok",
        "scenario_steps": [{"description": "s1", "method": "POST", "path": "/api/login",
                            "expect_status": 200, "json": {"username": "admin"}}],
    }]
    result = generate_framework(tmp_path, cases, FakeCfg(), {"endpoints": []})
    assert set(result["modules"]) == {"AUTH"}
    assert result["node_map"]["DEMO-AUTH-001"] == \
        "testcases/test_auth.py::test_auth[DEMO-AUTH-001]"
    for name in ("pytest.ini", "requirements.txt", "config.yaml", "conftest.py",
                 "common/api_client.py", "common/assertions.py", "common/runner.py",
                 "testcases/test_auth.py", "data/auth.yaml"):
        assert (tmp_path / name).exists(), name
    data = (tmp_path / "data" / "auth.yaml").read_text(encoding="utf-8")
    assert "DEMO-AUTH-003" in data and "scenario_steps" in data
    # 场景用例进入数据文件
    write_data_files(tmp_path, cases)
    assert "DEMO-AUTH-001" in (tmp_path / "data" / "auth.yaml").read_text(encoding="utf-8")


def test_generate_framework_py_files_compile(tmp_path: Path):
    import py_compile

    class FakeService:
        base_url = "http://x"
        timeout = 5
        test_account = {}
        login = {}

    class FakeCfg:
        project_code = "DEMO"
        service = FakeService()

    cases = _sample_cases()
    generate_framework(tmp_path, cases, FakeCfg(), {"endpoints": []})
    for py_file in tmp_path.rglob("*.py"):
        py_compile.compile(str(py_file), doraise=True)
